import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    if Path(csv_path).suffix.lower() != ".csv":
        raise ValueError("Входной файл должен быть CSV")
    if Path(xlsx_path).suffix.lower() != ".xlsx":
        raise ValueError("Выходной файл должен быть XLSX")

    with open(csv_path, "r", encoding="utf-8", newline="") as file:
        rows = list(csv.reader(file))

    if not rows or not rows[0] or any(not key.strip() for key in rows[0]):
        raise ValueError("CSV пустой или отсутствует заголовок")
    if len(set(rows[0])) != len(rows[0]):
        raise ValueError("Названия колонок CSV не должны повторяться")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    for column in ws.columns:
        max_length = 0

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = max(8, max_length + 2)

    Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
